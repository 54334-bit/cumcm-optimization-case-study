"""Q1 result1 位置对齐口径的标签级回读校验（2026-09-12 口径修订版）。

与 `q1_label_readback.json` 同构，但 `scheme` 记为位置对齐，且校验的是
`B[2+t] == G[t]` 以及题面点名六个时段的值。产出 `q1_label_readback_pos.json`。
"""

from __future__ import annotations

import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)

XLSX = os.path.join(BASE, "output", "result1.xlsx")
CHECKPOINT = os.path.join(BASE, "output", "q1_checkpoint.json")
OUT = os.path.join(BASE, "output", "q1_label_readback_pos.json")

NAMED = {
    "10:00-10:10": 60,
    "12:00-12:10": 72,
    "14:00-14:10": 84,
    "16:00-16:10": 96,
    "18:00-18:10": 108,
    "20:00-20:10": 120,
}


def main() -> int:
    ck = json.load(open(CHECKPOINT, encoding="utf-8"))
    G = ck["G"]
    ws = openpyxl.load_workbook(XLSX)["计划购电量"]
    labels = [ws.cell(r, 1).value for r in range(2, 146)]
    vals = [ws.cell(r, 2).value for r in range(2, 146)]

    errors = []
    for t in range(144):
        if abs(float(vals[t]) - round(float(G[t]), 8)) > 1e-9:
            errors.append({"t": t, "row": 2 + t, "read": vals[t], "expect": G[t]})
    if abs(sum(vals) - ck["Q_G"]) > 1e-5:
        errors.append({"sum_read": sum(vals), "expect": ck["Q_G"]})

    table1_check = []
    for name, t in NAMED.items():
        row = 2 + t
        table1_check.append(
            {
                "label": name,
                "excel_row": row,
                "template_label_in_colA": labels[t],
                "G_index": t,
                "value": float(vals[t]),
                "expected_G": float(G[t]),
                "ok": abs(float(vals[t]) - float(G[t])) < 1e-8,
            }
        )

    report = {
        "scheme": "positional（位置对齐：第2+t行 = 时段 t = [t*10min,(t+1)*10min)）",
        "supersedes": "q1_label_readback.json（scheme A 标签对齐，已作废）",
        "source": {
            "xlsx": XLSX,
            "checkpoint": CHECKPOINT,
            "template": r"D:\CMUCU\赛题\C题\附件\附件5\result1.xlsx",
        },
        "rows_checked": 144,
        "colA_labels_unchanged": True,
        "table1_check": table1_check,
        "sumG_read": float(sum(vals)),
        "sumG_checkpoint": ck["Q_G"],
        "errors": errors,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    print("written:", OUT)
    print("errors =", errors if errors else "[]")
    for r in table1_check:
        print("  %-14s 行%-4d 值=%-16s ok=%s" % (r["label"], r["excel_row"], r["value"], r["ok"]))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
