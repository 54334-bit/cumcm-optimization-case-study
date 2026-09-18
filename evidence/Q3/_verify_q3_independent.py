# -*- coding: utf-8 -*-
r"""Q3 交付**独立复算**（第二遍检查）：不 import 绘图/数据层代码，直接解提交件与逐段真值。

做法（刻意与 `q3_data_frozen.py` 不同源）：
  * 用 openpyxl 逐行读四张表（位置映射 col = 2 + t）；
  * 用 csv/json 直读逐段真值（G0/A/C/D/S/H）；
  * 自己实现读法 C 的四段计价、表3 前向填充、尾部集中度；
  * 与 `Q3交付\05_图件\q3_figures_manifest.json` 的 key_values 对表。
输出：每项「独立复算值 / manifest 值 / 差 / 判定」，末行总判。
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import numpy as np

Q3 = Path(r"D:\CMUCU\Q3交付")
XLSX = Q3 / "01_提交件" / "result3_v2.xlsx"
SEG = Path(r"D:\CMUCU\7.6对话\output\v2b_full\q76_v2b_full_seg.jsonl")
VOI = Path(r"D:\CMUCU\6对话\output\q3_voi_v2b\voi_v2b.json")
MAN = Q3 / "05_图件" / "q3_figures_manifest.json"
PRICE = Path(r"D:\CMUCU\B对话\clean\q1_clean.csv")
DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]


def read_price() -> np.ndarray:
    rows = list(csv.reader(PRICE.read_text(encoding="utf-8").splitlines()))
    hdr = [h.strip() for h in rows[0]]
    ci = [i for i, h in enumerate(hdr) if "price" in h.lower() or "电价" in h][0]
    return np.array([float(r[ci]) for r in rows[1:] if len(r) > ci], float)


def read_xlsx() -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(XLSX), read_only=True, data_only=True)
    out = {"sheets": wb.sheetnames}
    for nm in ("计划购电量", "调整购电量"):
        ws = wb[nm]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        out[nm + "_dates"] = [str(r[0])[:10] for r in rows]
        out[nm] = np.array([[float(v) for v in r[1:145]] for r in rows], float)
        out[nm + "_fee147"] = np.array([float(r[146]) for r in rows], float)
    ws = wb["充放电量"]
    blk, cur = {}, None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cur = str(r[0])[:10]
        if r[1] is None and r[2] is None:
            continue
        blk.setdefault(cur, []).append((str(r[1]),
                                        float(r[2]) if r[2] is not None else 0.0,
                                        float(r[3]) if r[3] is not None else 0.0))
    out["充放电量"] = blk
    ws = wb["紧急购电量"]
    segs, cur = [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cur = str(r[0])[:10]
        if r[1] is None and r[2] is None:
            continue
        segs.append((cur, str(r[1]), float(r[2])))
    out["紧急购电量"] = segs
    wb.close()
    return out


def read_seg() -> dict:
    recs = [json.loads(x) for x in SEG.read_text(encoding="utf-8").splitlines() if x.strip()]
    return {"dates": [r["date"] for r in recs],
            "G0": np.array([r["G0"] for r in recs], float),
            "A": np.array([r["A"] for r in recs], float),
            "C": np.array([r["C"] for r in recs], float),
            "D": np.array([r["D"] for r in recs], float),
            "H": np.array([r["H"] for r in recs], float),
            "S": np.array([r["S"] for r in recs], float)}


def cash_C(G0, A, H, p):
    p = p[None, :]
    j_plan = float((p * np.minimum(G0, A)).sum())
    j_cut = float((0.5 * p * np.maximum(G0 - A, 0)).sum())
    j_over = float((1.5 * p * np.maximum(A - G0, 0)).sum())
    j_emg = float((5.0 * p * H).sum())
    return j_plan, j_cut, j_over, j_emg, j_plan + j_cut + j_over + j_emg


def main() -> int:
    p = read_price()
    x = read_xlsx()
    s = read_seg()
    man = json.loads(MAN.read_text(encoding="utf-8"))
    kv = {f["name"]: f.get("key_values", {}) for f in man["figures"]}

    jp, jc, jo, je, jc_all = cash_C(s["G0"], s["A"], s["H"], p)
    rows = []
    rows.append(("J_plan", jp, kv["q3_02_cost_decomposition"]["v2_J_plan"]))
    rows.append(("J_cut(0.5p退款)", jc, kv["q3_02_cost_decomposition"]["v2_J_cut"]))
    rows.append(("J_over(1.5p超出)", jo, kv["q3_02_cost_decomposition"]["v2_J_over"]))
    rows.append(("J_emg(5p)", je, kv["q3_02_cost_decomposition"]["v2_J_emg"]))
    rows.append(("J_cash", jc_all, kv["q3_02_cost_decomposition"]["v2_J_cash"]))
    rows.append(("ΣH(kWh)", float(s["H"].sum()), kv["q3_02_cost_decomposition"]["v2_sumH_kWh"]))
    rows.append(("SOC min", float(s["S"].min()), kv["q3_05_soc_year"]["soc_min_kWh"]))
    rows.append(("SOC max", float(s["S"].max()), kv["q3_05_soc_year"]["soc_max_kWh"]))
    # 表3 四日期（前向填充）
    agg = {}
    cur = None
    for d, span, kwh in x["紧急购电量"]:
        if d:
            cur = d
        agg.setdefault(cur, []).append((span, kwh))
    for ds in DATES:
        n = len(agg.get(ds, []))
        tot = sum(k for _s, k in agg.get(ds, []))
        rows.append(("表3 %s 段数" % ds, n, kv["q3_04_paper_tables_graph"][ds]["n"]))
        rows.append(("表3 %s kWh" % ds, tot, kv["q3_04_paper_tables_graph"][ds]["kWh"]))
    # 表1/表2（提交件 vs 逐段真值）
    rows.append(("xlsx G0 vs seg 最大差",
                 float(np.abs(x["计划购电量"] - s["G0"]).max()), 0.0))
    rows.append(("xlsx A vs seg 最大差",
                 float(np.abs(x["调整购电量"] - s["A"]).max()), 0.0))
    blk_err = 0.0
    for di, ds in enumerate(s["dates"]):
        for k, (span, c_, d_) in enumerate(x["充放电量"].get(ds, [])[:6]):
            blk_err = max(blk_err, abs(c_ - s["C"][di, 24 * k:24 * (k + 1)].sum()),
                          abs(d_ - s["D"][di, 24 * k:24 * (k + 1)].sum()))
    rows.append(("表2 4h 块 vs 真值 最大差", blk_err, 0.0))
    # 表1/表2 全天购电费 合计（方案 A：本表购电量 × 该时段电价，4 位小数）
    rows.append(("计划表 全天费合计 vs Σp·G0",
                 float(x["计划购电量_fee147"].sum()), float((p * s["G0"]).sum())))
    rows.append(("调整表 全天费合计 vs Σp·A",
                 float(x["调整购电量_fee147"].sum()), float((p * s["A"]).sum())))
    # VOI 四臂
    if VOI.exists():
        d = json.loads(VOI.read_text(encoding="utf-8"))
        arms = sorted(d["arms"], key=lambda a: len(a["epochs"]))
        for a in arms:
            rows.append(("VOI %s J_cash" % (tuple(a["epochs"]),), a["totals"]["J_cash"],
                         kv["q3_03_rolling_value_voi"]["J_cash_yuan"][len(a["epochs"]) - 1]))
    bad = []
    print("=" * 92)
    for nm, mine, ref in rows:
        d_ = abs(float(mine) - float(ref))
        ok = d_ <= 0.01
        if not ok:
            bad.append((nm, mine, ref, d_))
        print("  %-26s 独立复算 %16.4f | manifest %16.4f | 差 %10.6f | %s"
              % (nm, mine, ref, d_, "OK" if ok else "**FAIL**"))
    print("=" * 92)
    print("检查 2（数值独立复算）：", "PASS" if not bad else "FAIL  %s" % bad)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
