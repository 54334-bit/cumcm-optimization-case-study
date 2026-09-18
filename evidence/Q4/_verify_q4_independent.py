# -*- coding: utf-8 -*-
r"""Q4 图件**独立复算**（第三遍检查）：不 import 绘图/数据层代码，直读证据与提交件重算。

做法（刻意与 `q4_data_frozen.py` / `q4_fig_main.py` 不同源）：
  * 自解析 `q4_cp2_checkpoint.json` 的 days[]，自己按 η=0.9 重建 SOC；
  * 自解析附件1 / 附件4 CSV，自己算分位与均值；
  * 自算 2×2 两个效应与分块闭合、四价格口径的相对偏差；
  * 自读 `result4-2.xlsx`（openpyxl）校验表1/表2/表3 与 checkpoint 的关系；
  * 与 `Q4交付\06_图件\q4_figures_manifest.json` 的 key_values 逐项对表。
输出：每项「独立复算 / manifest / 差 / 判定」，末行总判。
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import sys
from pathlib import Path

import numpy as np

Q4 = Path(r"D:\CMUCU\Q4交付")
EV = Q4 / "04_证据"
CKPT = EV / "q4_cp2_checkpoint.json"
CF2 = EV / "q4_counterfactual_2x2.json"
CF3 = EV / "q4_q43_2x2.json"
BLK = EV / "q4_price_block_decomp.json"
CAL = EV / "q4_probe_price_caliber_v2.json"
XLSX = Q4 / "01_提交件" / "result4-2.xlsx"
A4 = Path(r"D:\CMUCU\B对话\clean\attachment4_clean.csv")
A1 = Path(r"D:\CMUCU\B对话\clean\q1_clean.csv")
MAN = Path(r"D:\CMUCU\Q4交付\06_图件\q4_figures_manifest.json")
ETA = 0.9
T, N_DAYS, DAY0 = 144, 334, 31
DAYS7 = ["2025-03-20", "2025-05-15", "2025-06-21", "2025-07-15",
         "2025-09-23", "2025-10-15", "2025-12-21"]
# 审计 R10 发现 8：原 0.01 元只能挡量级错误；实测各差 ≤5e-5，收紧到 1e-3 元
TOL = 1e-3


def read_csv_matrix(path: Path, date_col: bool = True):
    rows = list(csv.reader(path.read_text(encoding="utf-8").splitlines()))
    hdr = [h.strip() for h in rows[0]]
    if date_col:
        dates = [r[0] for r in rows[1:]]
        arr = np.array([[float(v) for v in r[1:1 + T]] for r in rows[1:]], float)
        return hdr, dates, arr
    kept = [i for i, v in enumerate(rows[1]) if _isnum(v)]
    hdr2 = [hdr[i] for i in kept]
    arr = np.array([[float(r[i]) for i in kept] for r in rows[1:]], float)
    return hdr2, None, arr


def _isnum(v: str) -> bool:
    try:
        float(v)
        return True
    except (TypeError, ValueError):
        return False


def main() -> int:
    ck = json.loads(CKPT.read_text(encoding="utf-8"))
    days = ck["days"]
    cf2 = json.loads(CF2.read_text(encoding="utf-8"))
    cf3 = json.loads(CF3.read_text(encoding="utf-8"))
    blk = json.loads(BLK.read_text(encoding="utf-8"))
    cal = json.loads(CAL.read_text(encoding="utf-8"))
    man = json.loads(MAN.read_text(encoding="utf-8"))
    kv = {f["name"]: f.get("key_values", {}) for f in man["figures"]}

    rows: list[tuple[str, float, float]] = []
    rows.append(("checkpoint assertions_all_pass",
                 1.0 if ck.get("assertions_all_pass") else 0.0, 1.0))
    rows.append(("checkpoint 天数", float(len(days)), float(N_DAYS)))

    G = np.array([r["G"] for r in days], float)
    C = np.array([r["C"] for r in days], float)
    D = np.array([r["D"] for r in days], float)
    H = np.array([r["H"] for r in days], float)
    S0 = np.array([r["S0"] for r in days], float)
    S1 = np.array([r["S1"] for r in days], float)
    Jp = np.array([r["J_plan"] for r in days], float)
    Je = np.array([r["J_emg"] for r in days], float)

    # ---- 1. Q4-2 头条（由 checkpoint 逐日重算）
    rows.append(("Q4-2 J_plan 合计", float(Jp.sum()),
                 kv["q4_04_price_four"]["I_yuan"] - float(Je.sum())))
    rows.append(("Q4-2 全年费用(计划+紧急)", float(Jp.sum() + Je.sum()),
                 kv["q4_04_price_four"]["I_yuan"]))

    # ---- 2. 图 2：计划购电量
    qd = G.sum(axis=1)
    rows.append(("图2 QG 合计(kWh)", float(G.sum()), kv["q4_02_plan_heatmap"]["QG_kWh"]))
    rows.append(("图2 最大日计划量(kWh)", float(qd.max()),
                 kv["q4_02_plan_heatmap"]["max_day_QG_kWh"]))
    dmax = days[int(np.argmax(qd))]["d"]
    rows.append(("图2 最大日命中(0/1)",
                 1.0 if (dt.date(2025, 1, 1) + dt.timedelta(days=int(dmax))).isoformat()
                 == kv["q4_02_plan_heatmap"]["argmax_date"] else 0.0, 1.0))

    # ---- 3. 图 3：SOC 逐段重建（自己实现）
    worst = 0.0
    for i in range(len(days)):
        s = np.concatenate([[S0[i]], S0[i] + np.cumsum(ETA * C[i] - D[i] / ETA)])
        worst = max(worst, abs(s[-1] - S1[i]))
    rows.append(("图3 SOC 重建末端最大残差(kWh)", worst, 0.0))
    idx = {r["d"]: k for k, r in enumerate(days)}
    for ds in DAYS7:
        dd = (dt.date.fromisoformat(ds) - dt.date(2025, 1, 1)).days
        i = idx[dd]
        s = np.concatenate([[S0[i]], S0[i] + np.cumsum(ETA * C[i] - D[i] / ETA)])
        e = kv["q4_03_soc_price_7d"][ds]
        rows.append(("%s SOC S0" % ds, float(S0[i]), e["S0"]))
        rows.append(("%s SOC S1" % ds, float(S1[i]), e["S1"]))
        rows.append(("%s SOC min" % ds, float(s[1:].min()), e["soc_min"]))
        rows.append(("%s SOC max" % ds, float(s[1:].max()), e["soc_max"]))

    # ---- 4. 图 1：价格（自读 CSV，自分位）
    _, a4_dates, a4 = read_csv_matrix(A4)
    q = np.quantile(a4, [0.05, 0.5, 0.95], axis=0)
    e1 = kv["q4_01_price_caliber"]
    rows.append(("图1 附件4 全年均值(365d)", float(a4.mean()), e1["att4_mean_365d"]))
    rows.append(("图1 附件4 全年最大(365d)", float(a4.max()), e1["att4_max_365d"]))
    rows.append(("图1 附件4 全年最小(365d)", float(a4.min()), e1["att4_min_365d"]))
    for k, lab in ((0, "p05"), (1, "p50"), (2, "p95")):
        rows.append(("图1 20:00 分位 %s" % lab, float(q[k][120]),
                     e1["p05_p50_p95_at_20h_365d"][k]))
    rows.append(("图1 窗口天数", float(a4[DAY0:365].shape[0]), float(e1["n_days_window"])))
    hdr1, _, a1 = read_csv_matrix(A1, date_col=False)
    ci_p = [i for i, h in enumerate(hdr1) if "price" in h.lower() or "电价" in h][0]
    rows.append(("图1 附件1 均价", float(a1[:, ci_p].mean()), e1["att1_mean"]))
    rows.append(("图1 附件1 最大价", float(a1[:, ci_p].max()), e1["att1_max"]))

    # ---- 5. 图 4：四价格口径（自算相对偏差 + 跨文件互证）
    I, II, III, IV = (float(cal[k]) for k in ("I", "II", "III", "IV"))
    e4 = kv["q4_04_price_four"]
    rows.append(("图4 I 主口径", I, e4["I_yuan"]))
    rows.append(("图4 II 上周同日价", II, e4["II_yuan"]))
    rows.append(("图4 III 典型日价", III, e4["III_yuan"]))
    rows.append(("图4 IV 同周4周价", IV, e4["IV_yuan"]))
    rows.append(("图4 dII%", 100 * (II - I) / I, e4["dII_pct"]))
    rows.append(("图4 dIII%", 100 * (III - I) / I, e4["dIII_pct"]))
    rows.append(("图4 dIV%", 100 * (IV - I) / I, e4["dIV_pct"]))
    rows.append(("图4 III == cf2 G1P4",
                 1.0 if abs(III - float(cf2["cells"]["G1P4"]["total"])) <= TOL else 0.0, 1.0))
    rows.append(("图4 C(附件1) == cf2 G1P1",
                 1.0 if abs(float(cal["C_att1"]) - float(cf2["cells"]["G1P1"]["total"])) <= TOL else 0.0, 1.0))

    # ---- 6. 图 5：2×2 自算两效应 + 闭合
    c2, c3 = cf2["cells"], cf3["cells"]
    pe2 = float(c2["G1P4"]["total"]) - float(c2["G1P1"]["total"])
    re2 = float(c2["G4P4"]["total"]) - float(c2["G1P4"]["total"])
    pe3 = float(c3["plan1_settle4"]["J_cash"]) - float(c3["plan1_settle1"]["J_cash"])
    re3 = float(c3["plan4_settle4"]["J_cash"]) - float(c3["plan1_settle4"]["J_cash"])
    e5 = kv["q4_05_cf_2x2"]
    rows.append(("图5 Q4-2 价格效应", pe2, e5["Q4_2_price_effect"]))
    rows.append(("图5 Q4-2 重优化效应", re2, e5["Q4_2_reopt_effect"]))
    rows.append(("图5 Q4-3 价格效应", pe3, e5["Q4_3_price_effect"]))
    rows.append(("图5 Q4-3 重优化效应", re3, e5["Q4_3_reopt_effect"]))
    rows.append(("图5 Q4-2 左上格 = Q2 基线",
                 1.0 if abs(float(c2["G1P1"]["total"]) - float(cal["C_att1"])) <= TOL else 0.0, 1.0))
    rows.append(("图5 Q4-3 左上格 = Q3 v2 基线",
                 1.0 if abs(float(c3["plan1_settle1"]["J_cash"]) - 13120194.0621) <= TOL else 0.0, 1.0))
    rows.append(("图5 Q4-2 右下格 = 图4 I",
                 1.0 if abs(float(c2["G4P4"]["total"]) - I) <= TOL else 0.0, 1.0))
    rows.append(("图5 ①+②(Q4-2) = 相对 Q2 差", pe2 + re2, float(cf2["total_diff"])))

    # ---- 7. 图 6：分块（自求和）
    blocks = blk["blocks"]
    for name in ("c11", "c14", "c44"):
        rows.append(("图6 Σ%s" % name, float(np.sum(blk[name])),
                     float(c2[{"c11": "G1P1", "c14": "G1P4", "c44": "G4P4"}[name]]["total"])))
    rows.append(("图6 价格效应合计", float(np.sum(blk["c14"]) - np.sum(blk["c11"])), pe2))
    rows.append(("图6 重优化合计", float(np.sum(blk["c44"]) - np.sum(blk["c14"])), re2))
    for k, b in enumerate(blocks):
        rows.append(("图6 %s 价格效应" % b, float(blk["c14"][k] - blk["c11"][k]),
                     kv["q4_06_effect_split"]["eff_price_yuan"][k]))
        rows.append(("图6 %s 重优化效应" % b, float(blk["c44"][k] - blk["c14"][k]),
                     kv["q4_06_effect_split"]["eff_reopt_yuan"][k]))

    # ---- 8. 提交件 result4-2.xlsx 与 checkpoint 对表
    import openpyxl
    wb = openpyxl.load_workbook(str(XLSX), read_only=True, data_only=True)
    ws = wb["计划购电量"]
    xg, xf = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        xg.append([float(v) for v in r[1:1 + T]])
        xf.append(float(r[146]))
    xg = np.array(xg, float)
    rows.append(("提交件 表1 vs checkpoint G 最大差", float(np.abs(xg - G).max()), 0.0))
    rows.append(("提交件 表1 全天购电费 Σ", float(np.sum(xf)), float(Jp.sum())))
    ws = wb["充放电量"]
    # 表 2 结构：每天 6 行（4 小时块），"时刻"列只在前两行写 0:00 / 24:00，
    # 且这两行分别携带当日 0:00（=S0）与 24:00（=S1）的储电量（实测 2026-09-13）。
    blk_err, soc_err, cur, bi = 0.0, 0.0, None, 0
    didx = {r["d"]: k for k, r in enumerate(days)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cur = (dt.date.fromisoformat(str(r[0])[:10]) - dt.date(2025, 1, 1)).days
            bi = 0
        else:
            bi += 1
        i = didx[cur]
        if r[4] is not None and r[5] is not None:
            tag = str(r[4])
            ref = S1[i] if tag.startswith("24") else S0[i]
            soc_err = max(soc_err, abs(float(r[5]) - ref))
        if 0 <= bi <= 5:
            blk_err = max(blk_err,
                          abs(float(r[2] or 0.0) - C[i, 24 * bi:24 * (bi + 1)].sum()),
                          abs(float(r[3] or 0.0) - D[i, 24 * bi:24 * (bi + 1)].sum()))
        else:
            blk_err = max(blk_err, 1.0)
    rows.append(("提交件 表2 4h块 vs checkpoint C/D 最大差", blk_err, 0.0))
    rows.append(("提交件 表2 日界 SOC vs checkpoint 最大差", soc_err, 0.0))
    ws = wb["紧急购电量"]
    nrow = 0
    for _r in ws.iter_rows(min_row=2, values_only=True):
        if _r[1] is None and _r[2] is None:
            continue
        nrow += 1
    segs = 0
    for i in range(N_DAYS):
        pos = H[i] > 1e-9
        segs += int(np.sum(pos[1:] & ~pos[:-1]) + (1 if pos[0] else 0))
    rows.append(("提交件 表3 行数(段数)", float(nrow), float(segs)))
    rows.append(("提交件 表3 有紧急购电天数",
                 float(int((H.sum(axis=1) > 1e-9).sum())), 238.0))
    rows.append(("提交件 ΣH(kWh)", float(H.sum()), 57204.245176))
    wb.close()

    bad = []
    print("=" * 104)
    for nm, mine, ref in rows:
        d_ = abs(float(mine) - float(ref))
        ok = d_ <= TOL
        if not ok:
            bad.append((nm, float(mine), float(ref), d_))
        print("  %-40s 独立复算 %16.4f | manifest %16.4f | 差 %10.6f | %s"
              % (nm, mine, ref, d_, "OK" if ok else "**FAIL**"))
    print("=" * 104)
    print("检查 3（Q4 图件数值独立复算）：%s"
          % ("PASS（%d/%d）" % (len(rows), len(rows)) if not bad else "FAIL  %s" % bad))
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
