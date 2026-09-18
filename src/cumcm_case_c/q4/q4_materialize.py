# -*- coding: utf-8 -*-
"""Q4-2 物化：把 q4_main_solve 的 checkpoint 写入 result4-2.xlsx。

只读：
  - 模板 D:\\CMUCU\\赛题\\C题\\附件\\附件5\\result4-2.xlsx
  - 原始附件4 D:\\CMUCU\\赛题\\C题\\附件\\附件4.xlsx
  - checkpoint D:\\CMUCU\\8对话\\output\\q4_cp2_checkpoint.json
只写：
  - D:\\CMUCU\\8对话\\output\\result4-2.xlsx
  - D:\\CMUCU\\8对话\\output\\q4_cp4_materialize.json

本脚本是机械物化，不含任何建模/算法决策。
"""
from __future__ import annotations

import hashlib
import json
import os
from datetime import date, datetime, timedelta

import numpy as np
import openpyxl

ROOT = r"D:\CMUCU\8对话"
TPL = r"D:\CMUCU\赛题\C题\附件\附件5\result4-2.xlsx"
A4 = r"D:\CMUCU\赛题\C题\附件\附件4.xlsx"
CP_JSON = os.path.join(ROOT, "output", "q4_cp2_checkpoint.json")
OUT_XLSX = os.path.join(ROOT, "output", "result4-2.xlsx")
OUT_REPORT = os.path.join(ROOT, "output", "q4_cp4_materialize.json")

T = 144
D0 = 31
N_DAYS = 334
DT = 1.0 / 6.0
BLOCKS = [
    ("0:00-4:00", 0, 24),
    ("4:00-8:00", 24, 48),
    ("8:00-12:00", 48, 72),
    ("12:00-16:00", 72, 96),
    ("16:00-20:00", 96, 120),
    ("20:00-24:00", 120, 144),
]


def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def date_of(day_index: int) -> date:
    return date(2025, 1, 1) + timedelta(days=day_index)


def hhmm(minutes: int) -> str:
    return "24:00" if minutes >= 1440 else "%d:%02d" % (minutes // 60, minutes % 60)


def load_p4() -> np.ndarray:
    """读原始附件4：365 x 144，日期从 2025-01-01 起。"""
    wb = openpyxl.load_workbook(A4, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    assert len(rows) == 366, "附件4 应为 1 行表头 + 365 天，实际 %d" % len(rows)
    arr = np.asarray([[float(v) for v in r[1:1 + T]] for r in rows[1:]], dtype=float)
    assert arr.shape == (365, T), "附件4 形状 %s != (365,144)" % (arr.shape,)
    return arr


def load_checkpoint() -> dict:
    with open(CP_JSON, "r", encoding="utf-8") as fh:
        return json.load(fh)


def clear_data_rows(ws, first_row: int = 2) -> None:
    n = ws.max_row - first_row + 1
    if n > 0:
        ws.delete_rows(first_row, n)


def main() -> int:
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    p4 = load_p4()
    cp = load_checkpoint()
    days = cp["days"]

    struct_fail = []
    if len(days) != N_DAYS:
        struct_fail.append("checkpoint days 数 %d != %d" % (len(days), N_DAYS))
    for i, rec in enumerate(days):
        if int(rec["d"]) != D0 + i:
            struct_fail.append("checkpoint days[%d].d=%s != %d" % (i, rec["d"], D0 + i))
        for key in ("G", "C", "D", "H", "S0", "S1"):
            if key not in rec:
                struct_fail.append("checkpoint days[%d] 缺键 %s" % (i, key))
        if "G" in rec and len(rec["G"]) != T:
            struct_fail.append("checkpoint days[%d].G 长度 %d != %d" % (i, len(rec["G"]), T))
    if struct_fail:
        raise RuntimeError("checkpoint 结构自检失败：" + "; ".join(struct_fail[:5]))

    wb = openpyxl.load_workbook(TPL)
    ws1 = wb["计划购电量"]
    ws2 = wb["充放电量"]
    ws3 = wb["紧急购电量"]

    # ---------- 表1：计划购电量（位置映射 col = 2 + t） ----------
    if ws1.max_row < 1 + N_DAYS:
        raise RuntimeError("模板表1 行数不足：%d" % ws1.max_row)
    for i, rec in enumerate(days):
        r = 2 + i
        d = int(rec["d"])
        G = np.asarray(rec["G"], dtype=float)
        for t in range(T):
            ws1.cell(r, 2 + t).value = round(float(G[t]), 8)
        ws1.cell(r, 146).value = round(float(G.sum()), 6)
        ws1.cell(r, 147).value = round(float((p4[d] * G).sum()), 4)
        dv = ws1.cell(r, 1).value
        if dv is None or (dv.date() if isinstance(dv, datetime) else dv) != date_of(d):
            struct_fail.append("表1 第 %d 行日期 %r != %s" % (r, dv, date_of(d)))

    # ---------- 表2：充放电量（执行层实际；日期只在块首行） ----------
    clear_data_rows(ws2, 2)
    for i, rec in enumerate(days):
        d = int(rec["d"])
        C = np.asarray(rec["C"], dtype=float)
        D = np.asarray(rec["D"], dtype=float)
        for k, (label, a, b) in enumerate(BLOCKS):
            r = 2 + 6 * i + k
            ws2.cell(r, 1).value = (
                datetime(date_of(d).year, date_of(d).month, date_of(d).day) if k == 0 else None
            )
            ws2.cell(r, 2).value = label
            ws2.cell(r, 3).value = round(float(C[a:b].sum()), 8)
            ws2.cell(r, 4).value = round(float(D[a:b].sum()), 8)
            ws2.cell(r, 5).value = None
            ws2.cell(r, 6).value = None
        ws2.cell(2 + 6 * i, 5).value = datetime(2025, 1, 1, 0, 0).time()
        ws2.cell(2 + 6 * i, 6).value = round(float(rec["S0"]), 6)
        ws2.cell(3 + 6 * i, 5).value = "24:00"
        ws2.cell(3 + 6 * i, 6).value = round(float(rec["S1"]), 6)

    # ---------- 表3：紧急购电量（同日相邻正区间合并、跨日不合并） ----------
    clear_data_rows(ws3, 2)
    row = 2
    nseg = 0
    emg_days_written = 0
    emg_kwh_written = 0.0
    for i, rec in enumerate(days):
        d = int(rec["d"])
        H = np.asarray(rec["H"], dtype=float)
        t = 0
        segs = []
        while t < T:
            if H[t] > 1e-9:
                t0 = t
                q = 0.0
                while t < T and H[t] > 1e-9:
                    q += float(H[t])
                    t += 1
                segs.append((t0, t, q))
            else:
                t += 1
        for j, (t0, t1, q) in enumerate(segs):
            if j == 0:
                ws3.cell(row, 1).value = datetime(date_of(d).year, date_of(d).month, date_of(d).day)
            ws3.cell(row, 2).value = "%s-%s" % (hhmm(t0 * 10), hhmm(t1 * 10))
            ws3.cell(row, 3).value = round(q, 6)
            row += 1
            nseg += 1
            emg_kwh_written += q
        if segs:
            emg_days_written += 1

    wb.save(OUT_XLSX)
    wb.close()

    report = {
        "task": "q4_cp4_materialize",
        "generated_by": os.path.abspath(__file__),
        "template": TPL,
        "checkpoint": CP_JSON,
        "out_xlsx": OUT_XLSX,
        "n_days": len(days),
        "sheet1": {"rows": len(days), "col_mapping": "col = 2 + t; col2 = [0:00,0:10)"},
        "sheet2": {"block_rows": 6 * len(days), "blocks_per_day": 6},
        "sheet3": {
            "rows_written": nseg,
            "days_with_emergency": emg_days_written,
            "kwh_written": round(emg_kwh_written, 6),
            "merge_rule": "同日相邻 H>1e-9 区间合并；跨日不合并",
        },
        "struct_check_fail": struct_fail,
        "sha256": {
            "template": sha256_file(TPL),
            "checkpoint": sha256_file(CP_JSON),
            "out_xlsx": sha256_file(OUT_XLSX),
        },
    }
    with open(OUT_REPORT, "w", encoding="utf-8") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=1)

    print("written:", OUT_XLSX)
    print("report :", OUT_REPORT)
    print("表1 行数=%d | 表2 行数=%d | 表3 真实紧急段=%d 行 / %d 天 / %.6f kWh"
          % (len(days), 6 * len(days), nseg, emg_days_written, emg_kwh_written))
    print("结构自检失败:", struct_fail if struct_fail else "无")
    return 0 if not struct_fail else 1


if __name__ == "__main__":
    raise SystemExit(main())
