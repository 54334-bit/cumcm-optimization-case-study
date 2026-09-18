# -*- coding: utf-8 -*-
r"""Q3 图件线**唯一数据入口**（v2 交付口径：读法 C + 执行器 v2b）。

取数第一顺位（队长 2026-09-13 指定）：`D:\CMUCU\Q3交付\`；逐段明细取 `7.6对话\output\v2b_full\`。

口径：读法 C `p·min(G0,A) + 0.5p·(G0−A) + 1.5p·(A−G0) + 5p·H`；执行器 v2b；334 天；
表映射 `col = 2 + t`；头条 `J_cash = 13,120,194.06 元`、`ΣH = 42,672.2 kWh`；v1 = 13,369,682.34（仅对照）。
安全：只读；不重跑 `q76_v2b_full.py`；xlsx 字节 sha 不稳定 ⇒ 以单元格值验收（字节 sha 不符只告警）。

用法：$env:PYTHONPATH="D:\CMUCU\rag\.deps"; $env:PYTHONIOENCODING='utf-8'
      & $py "D:\CMUCU\6对话\code\q3_data_frozen.py"
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

import numpy as np

T, DT, N_DAYS = 144, 1.0 / 6.0, 334
D0 = dt.date(2025, 2, 1)
Q3 = Path(r"D:\CMUCU\Q3交付")
XLSX = Q3 / "01_提交件" / "result3_v2.xlsx"
XLSX_SHA_NOTE = "a9199e20d89060f81fd647690fee5d9614006173c79989c7c7dc1103c9918d10"   # 方案 A（4 位小数）版
V2B = Path(r"D:\CMUCU\7.6对话\output\v2b_full")
SEG_V2B = V2B / "q76_v2b_full_seg.jsonl"
SEG_E1 = V2B / "q76_v2b_full_e1_seg.jsonl"
DAILY = V2B / "q76_v2b_full_daily.csv"
VOI = Path(r"D:\CMUCU\6对话\output\q3_voi_v2b\voi_v2b.json")

EXPECT = {
    "J_cash_yuan": 13120194.06,
    "J_plan_yuan": 12381239.55,
    "J_adj_yuan": 542052.30,
    "J_emg_yuan": 196902.21,
    "sum_H_kWh": 42672.2,
    "v1_J_cash_yuan": 13369682.34,
    "n_days": 334,
    "S0_kWh": 6000.0,
    "soc_lo_kWh": 1200.0,
    "soc_hi_kWh": 10800.0,
}
# 容差：计数类零容差；交付公布值有舍入的按公布位数给容差
TOL = {"n_days": 0, "sum_H_kWh": 0.05, "S0_kWh": 1e-6}
PAPER_DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]


def sha256_of(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def load_submission(*, verbose: bool = True) -> dict:
    """按**单元格值**读四张表（列映射 col = 2 + t；主表第 147 列 = 全天购电费
    = **本表购电量 × 该时段电价**，方案 A / 4 位小数；全年总费用不落在表内）。"""
    import openpyxl

    wb = openpyxl.load_workbook(str(XLSX), read_only=True, data_only=True)
    want = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    if wb.sheetnames != want:
        raise ValueError("表名/顺序与模板不一致：%s" % wb.sheetnames)
    out: dict = {}
    for nm in want[:2]:
        ws = wb[nm]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        dates = [str(r[0])[:10] for r in rows]
        arr = np.zeros((len(rows), T), float)
        dayfee = np.zeros(len(rows), float)
        for i, r in enumerate(rows):
            for t in range(T):
                arr[i, t] = float(r[1 + t])
            dayfee[i] = float(r[146])
        out[nm] = {"dates": dates, "arr": arr, "dayfee": dayfee}
    ws = wb["充放电量"]
    blocks, cur = {}, None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cur = str(r[0])[:10]
        if r[1] is None and r[2] is None:
            continue
        # 注：表2 同日内只在首/末行写「时刻/储电量」，其余行为空 → 不能让 float(None) 崩
        blocks.setdefault(cur, []).append({
            "span": str(r[1]),
            "C": float(r[2]) if r[2] is not None else 0.0,
            "D": float(r[3]) if r[3] is not None else 0.0,
            "soc": float(r[5]) if r[5] is not None else None,
            "t": str(r[4]) if r[4] is not None else None})
    ws = wb["紧急购电量"]
    segs, cur = [], None
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            cur = str(r[0])[:10]
        if r[1] is None and r[2] is None:
            continue
        segs.append({"date": cur, "span": str(r[1]), "kWh": float(r[2])})
    wb.close()
    out["充放电量"] = blocks
    out["紧急购电量"] = segs
    if verbose:
        got = sha256_of(XLSX)
        flag = "一致" if got == XLSX_SHA_NOTE else "**已变**（须知：字节 sha 不作准，看锚点）"
        print("  提交件字节 sha256 %s… 与记录%s" % (got[:16], flag))
    return out


def load_seg(path: Path = SEG_V2B, *, verbose: bool = True) -> dict:
    """读逐段真值 JSONL（334 行：G0/A/C/D/S/H/R_PV/R_G）。"""
    recs = [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]
    if len(recs) != N_DAYS:
        raise ValueError("%s 行数 %d != %d" % (path.name, len(recs), N_DAYS))
    keys = ("G0", "A", "C", "D", "H", "R_PV", "R_G")
    arr = {k: np.array([r[k] for r in recs], float) for k in keys}
    S = np.array([r["S"] for r in recs], float)
    out = {"dates": [str(r["date"]) for r in recs], "d": np.array([int(r["d"]) for r in recs]),
           "S": S, "source": str(path), **arr}
    for k in keys:
        if arr[k].shape != (N_DAYS, T):
            raise ValueError("%s 形状 %s != (%d,%d)" % (k, arr[k].shape, N_DAYS, T))
    if S.shape != (N_DAYS, T + 1):
        raise ValueError("S 形状 %s != (%d,%d)" % (S.shape, N_DAYS, T + 1))
    if verbose:
        print("  逐段真值 %s：%d 天 × %d 段" % (path.name, N_DAYS, T))
    return out


def load_prices(*, verbose: bool = True) -> np.ndarray:
    """附件1 电价（典型日，144 段），源 `B对话/clean/q1_clean.csv`。"""
    p = Path(r"D:\CMUCU\B对话\clean\q1_clean.csv")
    rows = list(csv.reader(p.read_text(encoding="utf-8").splitlines()))
    hdr = [h.strip() for h in rows[0]]
    idx = [i for i, h in enumerate(hdr) if "price" in h.lower() or "电价" in h]
    if not idx:
        raise ValueError("q1_clean.csv 找不到电价列：%s" % hdr[:8])
    col = idx[0]
    vals = np.array([float(r[col]) for r in rows[1:] if len(r) > col], float)
    if vals.size != T:
        raise ValueError("电价长度 %d != %d" % (vals.size, T))
    if verbose:
        print("  电价：%d 段，%.4f~%.4f 元/kWh" % (vals.size, vals.min(), vals.max()))
    return vals


def cash_split_C(G0: np.ndarray, A: np.ndarray, H: np.ndarray, price: np.ndarray) -> dict:
    """读法 C 逐段分解（台账口径；第 147 列另按方案 A = 本表购电量 × 电价核对）：
    p·min(G0,A) + 0.5p·(G0−A)⁺ + 1.5p·(A−G0)⁺ + 5p·H。"""
    p = price[None, :]
    j_plan = float((p * np.minimum(G0, A)).sum())
    j_cut = float((0.5 * p * np.maximum(G0 - A, 0.0)).sum())
    j_over = float((1.5 * p * np.maximum(A - G0, 0.0)).sum())
    j_emg = float((5.0 * p * H).sum())
    return {"J_plan": j_plan, "J_adj": j_cut + j_over, "J_emg": j_emg,
            "J_cash": j_plan + j_cut + j_over + j_emg, "J_cut": j_cut, "J_over": j_over}


def _spans(v, thr: float = 1e-9) -> list[tuple[int, int, float]]:
    out, start = [], None
    for t, x in enumerate(v):
        pos = x > thr
        if pos and start is None:
            start = t
        elif not pos and start is not None:
            out.append((start, t - 1, float(v[start:t].sum())))
            start = None
    if start is not None:
        out.append((start, len(v) - 1, float(v[start:].sum())))
    return out


def anchors(seg: dict, price: np.ndarray) -> dict:
    cs = cash_split_C(seg["G0"], seg["A"], seg["H"], price)
    return {"J_cash_yuan": cs["J_cash"], "J_plan_yuan": cs["J_plan"],
            "J_adj_yuan": cs["J_adj"], "J_emg_yuan": cs["J_emg"],
            "sum_H_kWh": float(seg["H"].sum()), "n_days": int(seg["G0"].shape[0]),
            "S0_kWh": float(seg["S"][0, 0]),
            "soc_lo_kWh": float(seg["S"].min()), "soc_hi_kWh": float(seg["S"].max())}


def check_anchors(got: dict, *, verbose: bool = True) -> list[str]:
    bad = []
    for k, exp in EXPECT.items():
        if k not in got:                      # 非现算项（如 v1 对照值）只登记，不判 FAIL
            if verbose:
                print("  %-14s （参考值，非现算）%16.4f" % (k, exp))
            continue
        g = got[k]
        d = abs(g - exp)
        ok = d <= TOL.get(k, 0.01)
        if not ok:
            bad.append("%s: 现算 %r vs 交付 %r（差 %r）" % (k, g, exp, g - exp))
        if verbose:
            print("  %-14s 现算 %16.4f | 交付 %16.4f | 差 %10.6f | %s"
                  % (k, g, exp, g - exp, "OK" if ok else "**FAIL**"))
    return bad


def cross_check_xlsx(sub: dict, seg: dict, price: np.ndarray, *, verbose: bool = True) -> list[str]:
    """提交件 ↔ 逐段真值 逐格核对（计划/调整/充放电/紧急/全天费）。"""
    bad = []
    ref = cash_split_C(seg["G0"], seg["A"], seg["H"], price)["J_cash"]
    for nm, key in (("计划购电量", "G0"), ("调整购电量", "A")):
        arr = sub[nm]["arr"]
        if arr.shape != seg[key].shape:
            bad.append("%s 形状 %s != %s" % (nm, arr.shape, seg[key].shape))
            continue
        md = float(np.abs(arr - seg[key]).max())
        # 方案 A：本表购电量 × 该时段电价（4 位小数，日舍入上界 5e-5 → 全年容差 0.02 元）
        ref_fee = float((np.asarray(price, dtype=float) * np.asarray(seg[key], dtype=float)).sum())
        fee_gap = abs(float(sub[nm]["dayfee"].sum()) - ref_fee)
        if verbose:
            print("  [xlsx] %s 逐格最大差 %.3g；全天购电费合计 vs Σp·本表量 差 %.6g 元"
                  % (nm, md, fee_gap))
        if md > 1e-6:
            bad.append("%s 逐格最大差 %.3g 超容差" % (nm, md))
        if fee_gap > 0.02:
            bad.append("%s 全天购电费合计与 Σp·本表量 差 %.6g 元" % (nm, fee_gap))
    blk = sub["充放电量"]
    mc = md_ = 0.0
    for di, ds in enumerate(seg["dates"]):
        rows = blk.get(ds, [])
        if len(rows) != 6:
            bad.append("%s 充放电段数 %d != 6" % (ds, len(rows)))
            continue
        for k in range(6):
            c0, c1 = 24 * k, 24 * (k + 1)
            mc = max(mc, abs(rows[k]["C"] - seg["C"][di, c0:c1].sum()))
            md_ = max(md_, abs(rows[k]["D"] - seg["D"][di, c0:c1].sum()))
    if verbose:
        print("  [xlsx] 充放电 4h 段 vs 真值：充电最大差 %.3g / 放电最大差 %.3g" % (mc, md_))
    if mc > 1e-6 or md_ > 1e-6:
        bad.append("充放电 4h 段与真值差 充电 %.3g / 放电 %.3g" % (mc, md_))
    n_x = len(sub["紧急购电量"])
    n_true = sum(len(_spans(seg["H"][i])) for i in range(N_DAYS))
    if verbose:
        print("  [xlsx] 表3 段数 %d / 真值合并段数 %d" % (n_x, n_true))
    if n_x != n_true:
        bad.append("表3 段数 %d != %d" % (n_x, n_true))
    return bad


def paper_tables(sub: dict, seg: dict) -> dict:
    """论文表 1/2/3 的四日期数值（由提交件按位置映射现算）。"""
    out = {}
    slots_lab = (("10:00-10:10", 60), ("12:00-12:10", 72), ("14:00-14:10", 84),
                 ("16:00-16:10", 96), ("18:00-18:10", 108), ("20:00-20:10", 120))
    for ds in PAPER_DATES:
        i = sub["计划购电量"]["dates"].index(ds)
        out[ds] = {
            "slots": {lab: {"G0": round(float(sub["计划购电量"]["arr"][i, t]), 4),
                            "A": round(float(sub["调整购电量"]["arr"][i, t]), 4)}
                      for lab, t in slots_lab},
            "day_G0_kWh": round(float(sub["计划购电量"]["arr"][i].sum()), 4),
            "day_A_kWh": round(float(sub["调整购电量"]["arr"][i].sum()), 4),
            "day_fee_yuan": round(float(sub["计划购电量"]["dayfee"][i]), 4),
            "blocks": sub["充放电量"].get(ds, []),
            "emg": [s for s in sub["紧急购电量"] if s["date"] == ds],
        }
    return out


def daily_agg(seg: dict, price: np.ndarray) -> dict:
    """逐日聚合（画日序列/月度用）。"""
    p = price[None, :]
    j_plan = (p * np.minimum(seg["G0"], seg["A"])).sum(axis=1)
    j_cut = (0.5 * p * np.maximum(seg["G0"] - seg["A"], 0.0)).sum(axis=1)
    j_over = (1.5 * p * np.maximum(seg["A"] - seg["G0"], 0.0)).sum(axis=1)
    j_emg = (5.0 * p * seg["H"]).sum(axis=1)
    return {"date": seg["dates"], "J_plan": j_plan, "J_cut": j_cut, "J_over": j_over,
            "J_adj": j_cut + j_over, "J_emg": j_emg, "J_cash": j_plan + j_cut + j_over + j_emg,
            "QH": seg["H"].sum(axis=1), "QG0": seg["G0"].sum(axis=1), "QA": seg["A"].sum(axis=1),
            "QC": seg["C"].sum(axis=1), "QD": seg["D"].sum(axis=1),
            "soc_min": seg["S"][:, 1:].min(axis=1), "soc_max": seg["S"][:, 1:].max(axis=1),
            "R_PV": seg["R_PV"].sum(axis=1), "R_G": seg["R_G"].sum(axis=1)}


def main() -> int:
    print("== 提交件 ==")
    sub = load_submission()
    print("== 逐段真值 ==")
    seg = load_seg(SEG_V2B)
    price = load_prices()
    print("== 锚点对表（逐段真值现算 vs Q3交付 公布值）==")
    bad = check_anchors(anchors(seg, price))
    print("== 交叉核对（提交件 ↔ 逐段真值）==")
    bad += cross_check_xlsx(sub, seg, price)
    print("== 论文表（四日期）==")
    pt = paper_tables(sub, seg)
    for ds in PAPER_DATES:
        r = pt[ds]
        print("  %s 全天 G0=%.4f / A=%.4f / 费=%.4f 元；紧急段 %d（%s）"
              % (ds, r["day_G0_kWh"], r["day_A_kWh"], r["day_fee_yuan"], len(r["emg"]),
                 "；".join("%s %.4f" % (s["span"], s["kWh"]) for s in r["emg"]) or "无"))
    print("== v2b VOI（我方隔离重跑）==")
    if VOI.exists():
        d = json.loads(VOI.read_text(encoding="utf-8"))
        for a in d["arms"]:
            print("  %-14s J_cash=%15.2f  ΣH=%9.1f kWh"
                  % (str(tuple(a["epochs"])), a["totals"]["J_cash"], a["totals"]["QH"]))
        print("  self_check:", d["self_check"])
    print("\nANCHORS:", "PASS" if not bad else "FAIL")
    for b in bad:
        print("  !", b)
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
