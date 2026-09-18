# -*- coding: utf-8 -*-
"""整改验证：充放电量表的时段块是否已变成 4 小时（24 段）聚合，且 147 列不变。"""
import json
from openpyxl import load_workbook

XLSX = r"D:\CMUCU\7.6对话\交付\Q3交付\01_提交件\result3.xlsx"
SEG = r"D:\CMUCU\7.6对话\output\e4\q08_m0_seg.jsonl"

seg = [json.loads(l) for l in open(SEG, encoding="utf-8") if l.strip()]
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["充放电量"]
rows = [[c for c in r] for r in ws.iter_rows(min_row=2, values_only=True)]
print("充放电表数据行:", len(rows), "| 第1行:", rows[0][:6])

# 每天 6 块；第 d 天第 b 块 = 真值 C[t] 在 [24b, 24b+24) 的和
bad24 = bad36 = 0
maxdiff24 = 0.0
zero_blocks = 0
for d, day in enumerate(seg):
    C = day["C"]
    for b in range(6):
        got = rows[d * 6 + b][2] or 0.0
        exp24 = sum(C[24 * b:24 * b + 24])
        if abs(got - exp24) > 1e-6:
            bad24 += 1
            maxdiff24 = max(maxdiff24, abs(got - exp24))
        if abs(got) < 1e-12:
            zero_blocks += 1
print(f"与 24 段(4h) 窗口不符的块数 = {bad24}/2004  maxdiff={maxdiff24:.3e}")
print(f"恒为零的块数 = {zero_blocks}/2004")

ws1 = wb["计划购电量"]
tot = sum((r[146] or 0.0) for r in ws1.iter_rows(min_row=2, values_only=True))
# 方案 A（2026-09-13 统一）：第 147 列 = 本表购电量 × 该时段电价（4 位小数）
print(f"计划购电量表第147列全年合计 = {tot:.4f}（方案 A 目标 13077598.4544）")
