"""独立复算 Q3 工作簿的购电量与购电费。

原始附件、官方模板及其填报结果均不随本仓库再分发。使用者在已获许可的本地
副本上显式传入路径；脚本不会读取任何硬编码的个人目录。
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="独立复算 Q3 结果工作簿的费用与电量")
    parser.add_argument("--result-workbook", type=Path, required=True, help="本地 Q3 结果工作簿")
    parser.add_argument("--price-workbook", type=Path, required=True, help="本地附件 1 电价工作簿")
    parser.add_argument("--output", type=Path, required=True, help="输出审计文本路径")
    return parser


def _require_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise FileNotFoundError(f"{label}不存在或不是普通文件：{path}")


def _load_plan_sheet(workbook, name: str):
    import numpy as np

    rows = list(workbook[name].iter_rows(values_only=True))[1:]
    dates, values, daily_costs = [], [], []
    for row in rows:
        if row[0] is None:
            continue
        dates.append(row[0].date())
        values.append([float(value) for value in row[1:145]])
        daily_costs.append(float(row[146]))
    return dates, np.array(values), np.array(daily_costs)


def audit(result_workbook: Path, price_workbook: Path) -> str:
    """使用读法 C 和右端点电价复算，并返回 UTF-8 报告正文。"""
    import numpy as np
    import openpyxl

    _require_file(result_workbook, "结果工作簿")
    _require_file(price_workbook, "电价工作簿")
    price_book = openpyxl.load_workbook(price_workbook, read_only=True, data_only=True)
    prices = np.array([
        row[1] for row in price_book["Sheet1"].iter_rows(min_row=2, values_only=True)
        if row[1] is not None
    ], dtype=float)
    if prices.size != 144:
        raise ValueError(f"附件 1 预期含 144 个电价，实际为 {prices.size}")

    result_book = openpyxl.load_workbook(result_workbook, read_only=True, data_only=True)
    dates, planned, table_cost = _load_plan_sheet(result_book, "计划购电量")
    _, actual, _ = _load_plan_sheet(result_book, "调整购电量")
    index_by_date = {date: index for index, date in enumerate(dates)}
    emergency = np.zeros((len(dates), 144))
    current_date, segments, emergency_days = None, 0, set()
    for row in list(result_book["紧急购电量"].iter_rows(values_only=True))[1:]:
        if row[0] is not None:
            current_date = row[0].date()
        if row[2] is None or current_date not in index_by_date:
            continue
        match = re.fullmatch(r"(\d+):(\d+)-(\d+):(\d+)", str(row[1]))
        if match is None:
            raise ValueError(f"无法解析紧急购电时段：{row[1]!r}")
        start = int(match.group(1)) * 60 + int(match.group(2))
        end = int(match.group(3)) * 60 + int(match.group(4))
        slots = [slot for slot in range(144) if start < (slot + 1) * 10 <= end]
        if not slots:
            continue
        segments += 1
        emergency_days.add(current_date)
        emergency[index_by_date[current_date], slots] += float(row[2]) / len(slots)

    right_price = prices[None, :]
    fee_actual = float((actual * right_price).sum())
    fee_adjustment = float((0.5 * np.abs(planned - actual) * right_price).sum())
    fee_emergency = float((5.0 * emergency * right_price).sum())
    fee_recalculated = fee_actual + fee_adjustment + fee_emergency
    fee_table = float(table_cost.sum())
    volume_planned = float(planned.sum())
    volume_actual = float(actual.sum())
    volume_emergency = float(emergency.sum())
    difference = fee_recalculated - fee_table

    lines = [
        "Q3 结果工作簿全年购电量 / 购电费独立复算",
        "评价期和工作簿结构以使用者获准访问的本地官方材料为准。",
        "费用口径：读法 C，右端点电价；J = Σ[p·A + 0.5p·|G-A|] + Σ[5p·H]。",
        "",
        f"计划购电量 G = {volume_planned:,.6f} kWh",
        f"调整后购电量 A = {volume_actual:,.6f} kWh",
        f"紧急购电量 H = {volume_emergency:,.6f} kWh（{segments} 段/{len(emergency_days)} 天）",
        f"实际总购电量 A+H = {volume_actual + volume_emergency:,.6f} kWh",
        "",
        f"工作簿全天购电费列合计 = {fee_table:,.6f} 元",
        f"独立重算费用 = {fee_recalculated:,.6f} 元",
        f"差额（重算-表内） = {difference:,.4f} 元",
        f"执行电费 Σp·A = {fee_actual:,.2f} 元",
        f"调整电费 Σ0.5p·|G-A| = {fee_adjustment:,.2f} 元",
        f"紧急电费 Σ5p·H = {fee_emergency:,.2f} 元",
    ]
    return "\n".join(lines) + "\n"


def main() -> int:
    args = build_parser().parse_args()
    report = audit(args.result_workbook, args.price_workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(report, encoding="utf-8")
    print(f"已写入：{args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
